"""
GoldWatch Agent - DOCX Report Generator

Responsibilities:
- Create a readable daily gold report.
- Show available historical information.
- Clearly identify unavailable comparisons/forecasts.
- Include source and timestamps.

This module does NOT:
- Send emails.
- Read Gmail.
- Make financial decisions.
- Invent missing data.
"""

import sys
from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Pt


# ---------------------------------------------------------
# PROJECT ROOT
# ---------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent

# Allow imports from the GoldWatch project root.
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))


from database import get_latest_price
from comparison import get_comparison
from forecast import generate_forecast


REPORTS_DIR = BASE_DIR / "reports"


def format_currency(value):
    """
    Format an INR value.
    """

    return f"₹{value:,.2f}"


def add_title(document):
    """
    Add report title.
    """

    title = document.add_heading(
        "GOLDWATCH DAILY REPORT",
        level=0,
    )

    title.alignment = 1

    date_text = datetime.now().strftime(
        "%d %B %Y"
    )

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        f"Report generated: {date_text}"
    )

    run.bold = True


def add_current_prices(document):
    """
    Add latest available gold prices.
    """

    document.add_heading(
        "Latest Available Gold Price",
        level=1,
    )

    price_24k = get_latest_price("24K")
    price_22k = get_latest_price("22K")

    if not price_24k or not price_22k:

        document.add_paragraph(
            "Gold-price data is currently unavailable."
        )

        return

    table = document.add_table(
        rows=1,
        cols=3,
    )

    table.style = "Table Grid"

    header = table.rows[0].cells

    header[0].text = "Purity"
    header[1].text = "Price / gram"
    header[2].text = "Observation Date"

    for purity, row in [
        ("24K", price_24k),
        ("22K", price_22k),
    ]:

        cells = table.add_row().cells

        cells[0].text = purity

        cells[1].text = format_currency(
            row["price_per_gram"]
        )

        cells[2].text = row[
            "observation_date"
        ]

    document.add_paragraph(
        "Price unit: INR per gram."
    )

    document.add_paragraph(
        f"Source: {price_24k['source']}"
    )

    document.add_paragraph(
        f"Source timestamp: "
        f"{price_24k['api_timestamp']}"
    )

    document.add_paragraph(
        f"Retrieved by GoldWatch: "
        f"{price_24k['recorded_at']}"
    )


def add_comparison(document):
    """
    Add comparison when two different observation
    dates exist.
    """

    document.add_heading(
        "Price Comparison",
        level=1,
    )

    for purity in ["24K", "22K"]:

        result = get_comparison(purity)

        document.add_heading(
            purity,
            level=2,
        )

        if result["status"] != "success":

            document.add_paragraph(
                "Comparison unavailable: "
                + result["message"]
            )

            continue

        document.add_paragraph(
            f"Previous ({result['previous_observation_date']}): "
            f"{format_currency(result['previous_price'])}/gram"
        )

        document.add_paragraph(
            f"Current ({result['current_observation_date']}): "
            f"{format_currency(result['current_price'])}/gram"
        )

        document.add_paragraph(
            f"Change: "
            f"{format_currency(result['price_change'])}/gram "
            f"({result['percentage_change']:+.2f}%)"
        )

        document.add_paragraph(
            f"Direction: "
            f"{result['direction'].capitalize()}"
        )


def add_forecast(document):
    """
    Add forecast when enough historical data exists.
    """

    document.add_heading(
        "Tomorrow's Outlook",
        level=1,
    )

    for purity in ["24K", "22K"]:

        result = generate_forecast(purity)

        document.add_heading(
            purity,
            level=2,
        )

        if result["status"] != "success":

            document.add_paragraph(
                "Forecast unavailable: "
                + result["message"]
            )

            continue

        document.add_paragraph(
            f"Forecast range: "
            f"{format_currency(result['forecast_low'])}"
            f" – "
            f"{format_currency(result['forecast_high'])}"
            f" per gram"
        )

        document.add_paragraph(
            f"Forecast midpoint: "
            f"{format_currency(result['forecast_midpoint'])}"
            f" per gram"
        )

        document.add_paragraph(
            f"Expected movement: "
            f"{result['expected_change_percentage']:+.2f}%"
        )

        document.add_paragraph(
            f"Expected direction: "
            f"{result['direction'].capitalize()}"
        )

    document.add_paragraph(
        "Forecast caution: Any forecast is a statistical "
        "estimate and is not guaranteed. Actual gold prices "
        "may differ because of market conditions, currency "
        "movements, premiums, taxes and other factors."
    )


def add_footer(document):
    """
    Add report closing.
    """

    document.add_paragraph()

    document.add_paragraph(
        "Regards,"
    )

    document.add_paragraph(
        "Surya"
    )

    document.add_paragraph(
        "GoldWatch Agent"
    )


def generate_report():
    """
    Generate the GoldWatch DOCX report.
    """

    REPORTS_DIR.mkdir(
        exist_ok=True
    )

    document = Document()

    styles = document.styles

    styles["Normal"].font.size = Pt(10)

    add_title(document)

    add_current_prices(document)

    add_comparison(document)

    add_forecast(document)

    add_footer(document)

    filename = (
        "GoldWatch_Report_"
        + datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
        + ".docx"
    )

    report_path = (
        REPORTS_DIR / filename
    )

    document.save(report_path)

    return report_path


def main():
    """
    Generate a report.
    """

    print()
    print("=" * 70)
    print("GOLDWATCH - DOCX REPORT GENERATOR")
    print("=" * 70)

    try:

        report_path = generate_report()

        print()
        print("✅ Report generated successfully.")
        print()
        print(f"Report: {report_path}")

    except Exception as error:

        print()
        print("❌ Report generation failed.")
        print(f"Reason: {error}")

    print("=" * 70)


if __name__ == "__main__":
    main()